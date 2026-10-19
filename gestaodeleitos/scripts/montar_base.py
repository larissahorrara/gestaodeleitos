"""Monta a base de dados do painel de gestão de leitos no mesmo padrão do painel-sst-conformidade:
modelo estrela (d = dimensão, f = fato), uma Tabela do Excel por aba, abas LEIA-ME, Conferencia e Dicionario,
cópia em CSV (uma por tabela) e extração bruta em JSON. Todos os dados são fictícios.
Entrada: build/bases.json e build/calculos.json, gerados por scripts/extrair_dados.js a partir do index.html.
Saída: data/base-dados-leitos.xlsx, data/csv/ e data/dados.json
Uso (na raiz do repositório): python scripts/montar_base.py
"""
import csv, json, math, os
from datetime import datetime, date, timedelta
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

SAIDA = "data"
os.makedirs(f"{SAIDA}/csv", exist_ok=True)
HOJE = datetime(2026, 10, 19, 7, 0)
CALC = json.load(open("build/calculos.json", encoding="utf-8"))   # indicadores calculados pelo painel
BASES = json.load(open("build/bases.json", encoding="utf-8"))     # bases fictícias embutidas no painel

rd = lambda k: pd.DataFrame(BASES[k], dtype=str).fillna("")
leitos, censo, saidas, sol = rd("leitos"), rd("censo"), rd("saidas"), rd("solicitacoes")
atrasos, custos, aih = rd("atrasos"), rd("custos"), rd("receita_aih")
dt = lambda s: datetime.strptime(s, "%Y-%m-%d %H:%M") if s else None
d_ = lambda s: date.fromisoformat(s) if s else None
num = lambda s: float(s) if s not in ("", None) else None
inteiro = lambda s: int(float(s)) if s not in ("", None) else None

# ---------------- dimensões ----------------
ORDEM = list(dict.fromkeys(censo.sigla))
TIPO = {"CM": "Enfermaria clínica", "CC": "Enfermaria cirúrgica", "ORT": "Enfermaria cirúrgica", "UTI": "Terapia intensiva adulto",
        "UCI": "Cuidados intermediários", "PED": "Pediatria", "MAT": "Obstetrícia", "UTIN": "Terapia intensiva neonatal",
        "ONC": "Oncologia", "INF": "Isolamento"}
nome = dict(zip(censo.sigla, censo.unidade))
meta = censo.drop_duplicates("sigla").set_index("sigla")
diaria = custos.assign(v=custos.valor_diaria.astype(float)).groupby("sigla").v.sum()
aihv = dict(zip(aih.sigla, aih.receita_media_aih.astype(float)))
dUnidade = [{"id_unidade": s, "unidade": nome[s], "tipo_leito": TIPO[s], "leitos_instalados": int((leitos.sigla == s).sum()),
             "meta_taxa_ocupacao": float(meta.loc[s, "meta_taxa_ocupacao"]) / 100, "meta_media_permanencia": float(meta.loc[s, "meta_media_permanencia"]),
             "custo_diaria": round(float(diaria[s]), 2), "receita_media_aih": aihv[s]} for s in ORDEM]

dLeito = [{"id_leito": r.leito, "id_unidade": r.sigla, "ala": r.leito.split("-")[1][0], "numero": int(r.leito.split("-")[1][1:])} for r in leitos.itertuples()]

DIAS = ["segunda-feira", "terça-feira", "quarta-feira", "quinta-feira", "sexta-feira", "sábado", "domingo"]
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
dCalendario = []
for s in sorted(set(censo.data)):
    x = date.fromisoformat(s)
    dCalendario.append({"data": x, "dia": x.day, "mes": x.month, "nome_mes": MESES[x.month - 1], "ano": x.year,
                        "dia_semana": x.isoweekday(), "nome_dia": DIAS[x.weekday()], "fim_de_semana": "Sim" if x.weekday() >= 5 else "Não"})

procs = leitos[leitos.procedimento != ""].drop_duplicates(["sigla", "procedimento"]).sort_values(["sigla", "procedimento"])
procs = procs.set_index("sigla").loc[[s for s in ORDEM if s in set(procs.sigla)]].reset_index()
dProcedimento, id_proc = [], {}
for i, r in enumerate(procs.itertuples(), 1):
    pid = f"PR{i:02d}"; id_proc[(r.sigla, r.procedimento)] = pid
    dProcedimento.append({"id_procedimento": pid, "procedimento": r.procedimento, "id_unidade": r.sigla, "permanencia_ref_dias": int(float(r.permanencia_ref_dias))})

ACOES = {
    "Alta social": ("Serviço Social", "Triagem de risco social na admissão; plano de alta iniciado em até 48 h; articulação com CRAS/CREAS e família; reunião semanal de casos longos."),
    "Falta de transporte": ("Transporte sanitário / NIR", "Agendar o transporte na véspera da alta prevista; escala de ambulância para altas até 10 h; sala de alta para aguardar fora do leito."),
    "Falta de acompanhante ou responsável": ("Serviço Social", "Identificar o responsável na admissão; contato ativo com a família na véspera; protocolo de alta com rede de apoio."),
    "Aguardando leito de retaguarda": ("NIR / Regulação estadual", "Solicitar a vaga assim que a alta for previsível; pactuar cotas com hospitais de cuidados prolongados e ILPI; escalonar casos acima de 5 dias."),
    "Aguardando atenção domiciliar": ("Atenção domiciliar (SAD)", "Encaminhar ao SAD/Melhor em Casa no início da internação; checklist de oxigênio e equipamentos domiciliares."),
    "Aguardando insumo ou equipamento para alta": ("Farmácia e suprimentos", "Prever órteses, curativos e oxigênio na véspera; estoque mínimo para alta; dispensação antes das 10 h."),
    "Pendência judicial ou documental": ("Jurídico / Serviço Social", "Acionar Jurídico e Serviço Social no primeiro dia; modelo padronizado de comunicação ao Ministério Público e à Defensoria."),
    "Aguardando transferência para outro serviço": ("NIR / Regulação estadual", "Inserir o pedido na regulação junto com a previsão de alta; acompanhamento diário pelo NIR."),
}
dMotivoDesocupacao, id_mot = [], {}
for i, (m, (area, acao)) in enumerate(ACOES.items(), 1):
    id_mot[m] = f"MD{i:02d}"
    dMotivoDesocupacao.append({"id_motivo": f"MD{i:02d}", "motivo": m, "area_responsavel": area, "acao_recomendada": acao})

comps = custos.drop_duplicates("componente")
dComponenteCusto, id_comp = [], {}
for i, r in enumerate(comps.itertuples(), 1):
    id_comp[r.componente] = f"CC{i:02d}"
    dComponenteCusto.append({"id_componente": f"CC{i:02d}", "componente": r.componente, "fator_dia_de_espera": float(r.fator_dia_de_espera)})

dMetas = [
    {"indicador": "Taxa de ocupação (30 dias)", "meta": 0.85, "unidade_medida": "%", "sentido": "Faixa", "regra": "Atenção a partir de 85%; crítico a partir de 90%"},
    {"indicador": "Ocupação instantânea", "meta": 0.85, "unidade_medida": "%", "sentido": "Faixa", "regra": "Crítico a partir de 90% ou com três unidades nesse patamar"},
    {"indicador": "Média de permanência", "meta": round(CALC["kpi_periodo"][-1]["meta_mp"], 1), "unidade_medida": "dias", "sentido": "Menor é melhor", "regra": "Meta do hospital ponderada pelos pacientes-dia das metas das unidades"},
    {"indicador": "Pacientes em vermelho (Kanban)", "meta": 0.15, "unidade_medida": "% dos internados", "sentido": "Menor é melhor", "regra": "Atenção a partir de 15%; crítico a partir de 25%"},
    {"indicador": "Saídas até 10 h", "meta": 0.30, "unidade_medida": "% das saídas", "sentido": "Maior é melhor", "regra": "Crítico abaixo de 15%; atenção abaixo de 30%"},
    {"indicador": "Alta médica até a saída", "meta": 120, "unidade_medida": "minutos (mediana)", "sentido": "Menor é melhor", "regra": "Acima de 120 min indica alta tardia"},
    {"indicador": "Saída até início da higienização", "meta": 30, "unidade_medida": "minutos (mediana)", "sentido": "Menor é melhor", "regra": "Acima de 30 min indica demora em acionar a limpeza"},
    {"indicador": "Higienização concorrente", "meta": 60, "unidade_medida": "minutos", "sentido": "Menor é melhor", "regra": "Acima do alvo atrasa a liberação do leito"},
    {"indicador": "Higienização terminal", "meta": 120, "unidade_medida": "minutos", "sentido": "Menor é melhor", "regra": "Acima do alvo atrasa a liberação do leito"},
    {"indicador": "Espera por leito", "meta": 6, "unidade_medida": "horas", "sentido": "Menor é melhor", "regra": "Crítico com espera pendente acima de 24 h ou menos de 40% atendidos em até 6 h"},
    {"indicador": "Leitos-dia perdidos com desocupação", "meta": 0.04, "unidade_medida": "% dos leitos-dia", "sentido": "Menor é melhor", "regra": "Atenção a partir de 4%; crítico a partir de 8%"},
    {"indicador": "Saldo de leitos do dia", "meta": 2, "unidade_medida": "leitos", "sentido": "Maior é melhor", "regra": "Crítico com saldo negativo; atenção com saldo até 2"},
]

# ---------------- fatos (valores + fórmulas) ----------------
fCensoDiario = [{"data": d_(r.data), "id_unidade": r.sigla, "leitos_operacionais": int(r.leitos_operacionais), "leitos_bloqueados": int(r.leitos_bloqueados),
                 "pacientes_dia": int(r.pacientes_dia), "admissoes": int(r.admissoes), "altas": int(r.altas), "obitos": int(r.obitos),
                 "obitos_ate_24h": int(r.obitos_ate_24h), "transferencias_externas": int(r.transferencias_externas)} for r in censo.itertuples()]
for x in fCensoDiario:
    x["saidas"] = x["altas"] + x["obitos"] + x["transferencias_externas"]
    x["taxa_ocupacao"] = x["pacientes_dia"] / x["leitos_operacionais"]

ref_proc = {p["id_procedimento"]: p["permanencia_ref_dias"] for p in dProcedimento}
fLeitosCenso = []
for r in leitos.itertuples():
    pid = id_proc.get((r.sigla, r.procedimento), "")
    dias = inteiro(r.dias_internacao) if r.status == "Ocupado" else None
    refd = ref_proc.get(pid)
    x = {"id_leito": r.leito, "id_unidade": r.sigla, "status": r.status, "etapa_fluxo": r.etapa_kanban, "prontuario": r.paciente, "iniciais": r.iniciais,
         "idade": inteiro(r.idade), "sexo": r.sexo, "id_procedimento": pid, "data_internacao": dt(r.data_internacao), "dias_internacao": dias,
         "permanencia_ref_dias": refd, "razao_permanencia": (dias / refd) if dias is not None and refd else None,
         "classificacao_kanban": "" if dias is None else ("Verde" if dias <= refd / 2 else "Amarelo" if dias <= refd else "Vermelho"),
         "previsao_alta": d_(r.previsao_alta), "origem": r.origem, "pendencia": r.pendencia, "motivo_bloqueio": r.motivo_bloqueio,
         "bloqueado_desde": dt(r.bloqueado_desde), "tipo_higienizacao": r.tipo_higienizacao, "inicio_higienizacao": dt(r.inicio_higienizacao),
         "reservado_para": r.reservado_para, "reservado_desde": dt(r.reservado_desde), "vago_desde": dt(r.vago_desde)}
    assert x["classificacao_kanban"] == r.classificacao_kanban, (r.leito, x["classificacao_kanban"], r.classificacao_kanban)
    fLeitosCenso.append(x)

fSaidas = []
for i, r in enumerate(saidas.itertuples(), 1):
    sp = dt(r.saida_paciente)
    fSaidas.append({"id_saida": f"SA{i:04d}", "data": d_(r.data), "id_unidade": r.sigla, "tipo_saida": r.tipo_saida, "alta_medica": dt(r.alta_medica),
                    "saida_paciente": sp, "inicio_higienizacao": dt(r.inicio_higienizacao), "fim_higienizacao": dt(r.fim_higienizacao),
                    "tipo_higienizacao": r.tipo_higienizacao, "nova_ocupacao": dt(r.nova_ocupacao), "min_alta_ate_saida": int(r.min_alta_ate_saida),
                    "min_saida_ate_inicio_hig": int(r.min_saida_ate_inicio_hig), "min_higienizacao": int(r.min_higienizacao),
                    "min_liberacao_ate_ocupacao": int(r.min_liberacao_ate_ocupacao), "min_giro_total": int(r.min_giro_total),
                    "saida_ate_10h": "" if r.tipo_saida == "Óbito" else ("Sim" if sp.hour < 10 else "Não")})

fSolicitacoes = [{"id_solicitacao": r.id, "data_solicitacao": dt(r.data_solicitacao), "origem": r.origem, "id_unidade_destino": r.sigla_destino,
                  "prioridade": r.prioridade, "data_alocacao": dt(r.data_alocacao), "status": r.status, "tempo_espera_h": float(r.tempo_espera_h)} for r in sol.itertuples()]
for x in fSolicitacoes:
    x["atendida_ate_6h"] = "Sim" if x["status"] == "Atendida" and x["tempo_espera_h"] <= 6 else "Não"

fCustoDiaria = []
for r in custos.itertuples():
    fat = float(r.fator_dia_de_espera)
    fCustoDiaria.append({"id_unidade": r.sigla, "id_componente": id_comp[r.componente], "valor_diaria": float(r.valor_diaria),
                         "fator_dia_de_espera": fat, "valor_dia_espera": float(r.valor_diaria) * fat})

mp1 = {k["sigla"]: k["media_permanencia"] for k in CALC["kpi_periodo"]}   # permanência média (1 casa), como no painel
from decimal import Decimal, ROUND_HALF_UP   # arredonda como o ROUND do Excel
dia_espera = {s: float(sum(Decimal(str(c["valor_diaria"])) * Decimal(str(c["fator_dia_de_espera"])) for c in fCustoDiaria if c["id_unidade"] == s).quantize(Decimal("0.01"), ROUND_HALF_UP)) for s in ORDEM}
xround = lambda v, n=0: float(Decimal(str(v)).quantize(Decimal(1).scaleb(-n), ROUND_HALF_UP))
fAtrasosDesocupacao = []
for r in atrasos.itertuples():
    dias = float(r.dias_atraso); ce = dia_espera[r.sigla]
    fAtrasosDesocupacao.append({"id_atraso": r.id, "id_leito": r.leito, "id_unidade": r.sigla, "prontuario": r.paciente, "id_motivo": id_mot[r.motivo],
                                "data_alta_medica": dt(r.data_alta_medica), "data_desocupacao": dt(r.data_desocupacao), "situacao": r.situacao,
                                "dias_atraso": dias, "custo_dia_espera": ce, "custo_total": int(xround(dias * ce)),
                                "internacoes_perdidas": round(dias / mp1[r.sigla], 2), "receita_nao_realizada": round(dias / mp1[r.sigla] * aihv[r.sigla])})
# confere com o cálculo do painel
ep = {e["id"]: e for e in CALC["desoc_episodios"]}
for x in fAtrasosDesocupacao:
    e = ep[x["id_atraso"]]
    assert abs(e["custo_total"] - x["custo_total"]) <= 1 and abs(e["receita_nao_realizada"] - x["receita_nao_realizada"]) <= 1, x["id_atraso"]

# ---------------- definição das tabelas (colunas, formatos, fórmulas, dicionário) ----------------
D, I, F1, F2, P, DT, DTH, R = "dd/mm/yyyy", "#,##0", "#,##0.0", "#,##0.00", "0.0%", "dd/mm/yyyy", "dd/mm/yyyy hh:mm", '"R$" #,##0.00'
RS0 = '"R$" #,##0'
REL_U = "dUnidade[id_unidade]"
T = {}  # nome -> (linhas, [(coluna, formato, formula|None, tipo, descricao, relacionamento)])
T["dUnidade"] = (dUnidade, [
    ("id_unidade", None, None, "texto", "Sigla da unidade de internação (chave)", ""),
    ("unidade", None, None, "texto", "Nome da unidade", ""),
    ("tipo_leito", None, None, "texto", "Perfil assistencial do leito", ""),
    ("leitos_instalados", I, None, "inteiro", "Leitos instalados no censo de referência", ""),
    ("meta_taxa_ocupacao", P, None, "percentual", "Meta de taxa de ocupação da unidade", ""),
    ("meta_media_permanencia", F1, None, "decimal", "Meta de média de permanência, em dias", ""),
    ("custo_diaria", R, None, "moeda", "Custo padrão da diária (soma dos componentes de fCustoDiaria)", ""),
    ("receita_media_aih", R, None, "moeda", "Receita média por internação (AIH) usada para estimar a receita não realizada", "")])
T["dLeito"] = (dLeito, [
    ("id_leito", None, None, "texto", "Código do leito (unidade-ala+número)", ""),
    ("id_unidade", None, None, "texto", "Unidade do leito", REL_U),
    ("ala", None, None, "texto", "Ala da unidade (A ou B)", ""),
    ("numero", I, None, "inteiro", "Número do leito na ala", "")])
T["dCalendario"] = (dCalendario, [
    ("data", D, None, "data", "Dia do período de análise (chave)", ""),
    ("dia", I, None, "inteiro", "Dia do mês", ""), ("mes", I, None, "inteiro", "Mês (1 a 12)", ""),
    ("nome_mes", None, None, "texto", "Nome do mês", ""), ("ano", I, None, "inteiro", "Ano", ""),
    ("dia_semana", I, None, "inteiro", "Dia da semana (1 = segunda, 7 = domingo)", ""),
    ("nome_dia", None, None, "texto", "Nome do dia da semana", ""), ("fim_de_semana", None, None, "texto", "Sim para sábado e domingo", "")])
T["dProcedimento"] = (dProcedimento, [
    ("id_procedimento", None, None, "texto", "Código do procedimento (chave)", ""),
    ("procedimento", None, None, "texto", "Procedimento principal da internação", ""),
    ("id_unidade", None, None, "texto", "Unidade em que o procedimento é realizado", REL_U),
    ("permanencia_ref_dias", I, None, "inteiro", "Tempo de referência de permanência (padrão SIGTAP simplificado), base do Kanban", "")])
T["dMotivoDesocupacao"] = (dMotivoDesocupacao, [
    ("id_motivo", None, None, "texto", "Código do motivo de atraso na desocupação (chave)", ""),
    ("motivo", None, None, "texto", "Motivo não clínico que impede a saída após a alta médica", ""),
    ("area_responsavel", None, None, "texto", "Área responsável por destravar o caso", ""),
    ("acao_recomendada", None, None, "texto", "Prática recomendada de planejamento de alta", "")])
T["dComponenteCusto"] = (dComponenteCusto, [
    ("id_componente", None, None, "texto", "Código do componente de custo (chave)", ""),
    ("componente", None, None, "texto", "Componente da diária hospitalar", ""),
    ("fator_dia_de_espera", P, None, "percentual", "Parte do componente que continua sendo gasta enquanto o paciente de alta espera no leito", "")])
T["fCensoDiario"] = (fCensoDiario, [
    ("data", D, None, "data", "Dia do censo", "dCalendario[data]"),
    ("id_unidade", None, None, "texto", "Unidade", REL_U),
    ("leitos_operacionais", I, None, "inteiro", "Leitos instalados menos bloqueados no dia", ""),
    ("leitos_bloqueados", I, None, "inteiro", "Leitos fora de operação no dia", ""),
    ("pacientes_dia", I, None, "inteiro", "Pacientes internados no censo do dia (pacientes-dia)", ""),
    ("admissoes", I, None, "inteiro", "Entradas na unidade no dia", ""),
    ("altas", I, None, "inteiro", "Altas no dia", ""), ("obitos", I, None, "inteiro", "Óbitos no dia", ""),
    ("obitos_ate_24h", I, None, "inteiro", "Óbitos com menos de 24 h de internação (fora da mortalidade institucional)", ""),
    ("transferencias_externas", I, None, "inteiro", "Transferências para outros serviços no dia", ""),
    ("saidas", I, "=G{r}+H{r}+J{r}", "inteiro (fórmula)", "Saídas = altas + óbitos + transferências externas", ""),
    ("taxa_ocupacao", P, "=IF(C{r}=0,0,E{r}/C{r})", "percentual (fórmula)", "Taxa de ocupação do dia = pacientes-dia ÷ leitos operacionais", "")])
T["fLeitosCenso"] = (fLeitosCenso, [
    ("id_leito", None, None, "texto", "Leito", "dLeito[id_leito]"),
    ("id_unidade", None, None, "texto", "Unidade", REL_U),
    ("status", None, None, "texto", "Situação no censo: Ocupado, Vago, Em higienização, Reservado ou Bloqueado", ""),
    ("etapa_fluxo", None, None, "texto", "Etapa do fluxo do leito (coluna do Kanban)", ""),
    ("prontuario", None, None, "texto", "Prontuário fictício do paciente", ""),
    ("iniciais", None, None, "texto", "Iniciais fictícias do paciente", ""),
    ("idade", I, None, "inteiro", "Idade em anos", ""), ("sexo", None, None, "texto", "Sexo (F ou M)", ""),
    ("id_procedimento", None, None, "texto", "Procedimento da internação", "dProcedimento[id_procedimento]"),
    ("data_internacao", DTH, None, "data e hora", "Data e hora da internação", ""),
    ("dias_internacao", I, "=IF(C{r}=\"Ocupado\",INT('LEIA-ME'!$B$4-J{r}),\"\")", "inteiro (fórmula)", "Dias completos de internação até a data de referência", ""),
    ("permanencia_ref_dias", I, "=IF(I{r}=\"\",\"\",INDEX(dProcedimento!$D:$D,MATCH(I{r},dProcedimento!$A:$A,0)))", "inteiro (fórmula)", "Tempo de referência do procedimento", ""),
    ("razao_permanencia", F2, "=IF(K{r}=\"\",\"\",K{r}/L{r})", "decimal (fórmula)", "Dias de internação ÷ tempo de referência", ""),
    ("classificacao_kanban", None, "=IF(K{r}=\"\",\"\",IF(K{r}<=L{r}/2,\"Verde\",IF(K{r}<=L{r},\"Amarelo\",\"Vermelho\")))", "texto (fórmula)", "Verde até metade do tempo de referência; Amarelo até o tempo; Vermelho acima", ""),
    ("previsao_alta", D, None, "data", "Data prevista para a alta", ""),
    ("origem", None, None, "texto", "Porta de entrada do paciente", ""),
    ("pendencia", None, None, "texto", "Pendência que atrasa a alta ou a desocupação", ""),
    ("motivo_bloqueio", None, None, "texto", "Motivo do bloqueio do leito", ""),
    ("bloqueado_desde", DTH, None, "data e hora", "Início do bloqueio", ""),
    ("tipo_higienizacao", None, None, "texto", "Concorrente ou terminal", ""),
    ("inicio_higienizacao", DTH, None, "data e hora", "Início da higienização em andamento", ""),
    ("reservado_para", None, None, "texto", "Origem do paciente para quem o leito está reservado", ""),
    ("reservado_desde", DTH, None, "data e hora", "Início da reserva", ""),
    ("vago_desde", DTH, None, "data e hora", "Momento em que o leito ficou vago", "")])
T["fSaidas"] = (fSaidas, [
    ("id_saida", None, None, "texto", "Código da saída (chave)", ""),
    ("data", D, None, "data", "Dia da saída", "dCalendario[data]"),
    ("id_unidade", None, None, "texto", "Unidade", REL_U),
    ("tipo_saida", None, None, "texto", "Alta, Óbito ou Transferência externa", ""),
    ("alta_medica", DTH, None, "data e hora", "Registro da alta médica", ""),
    ("saida_paciente", DTH, None, "data e hora", "Saída física do paciente do leito", ""),
    ("inicio_higienizacao", DTH, None, "data e hora", "Início da limpeza do leito", ""),
    ("fim_higienizacao", DTH, None, "data e hora", "Fim da limpeza do leito", ""),
    ("tipo_higienizacao", None, None, "texto", "Concorrente ou terminal", ""),
    ("nova_ocupacao", DTH, None, "data e hora", "Entrada do próximo paciente no leito", ""),
    ("min_alta_ate_saida", I, None, "inteiro", "Minutos entre a alta médica e a saída", ""),
    ("min_saida_ate_inicio_hig", I, None, "inteiro", "Minutos entre a saída e o início da limpeza", ""),
    ("min_higienizacao", I, None, "inteiro", "Duração da higienização em minutos", ""),
    ("min_liberacao_ate_ocupacao", I, None, "inteiro", "Minutos entre o leito liberado e a nova ocupação", ""),
    ("min_giro_total", I, None, "inteiro", "Minutos da saída até a nova ocupação", ""),
    ("saida_ate_10h", None, "=IF(D{r}=\"Óbito\",\"\",IF(HOUR(F{r})<10,\"Sim\",\"Não\"))", "texto (fórmula)", "Sim quando a saída ocorreu antes das 10 h (óbitos fora)", "")])
T["fSolicitacoes"] = (fSolicitacoes, [
    ("id_solicitacao", None, None, "texto", "Código do pedido de leito (chave)", ""),
    ("data_solicitacao", DTH, None, "data e hora", "Registro do pedido no NIR", ""),
    ("origem", None, None, "texto", "Origem do pedido", ""),
    ("id_unidade_destino", None, None, "texto", "Unidade solicitada", REL_U),
    ("prioridade", None, None, "texto", "Emergência, Urgência ou Eletiva", ""),
    ("data_alocacao", DTH, None, "data e hora", "Alocação do leito (vazio se pendente ou cancelado)", ""),
    ("status", None, None, "texto", "Atendida, Cancelada ou Pendente", ""),
    ("tempo_espera_h", F1, None, "decimal", "Horas entre o pedido e a alocação (pendentes: até o censo)", ""),
    ("atendida_ate_6h", None, "=IF(AND(G{r}=\"Atendida\",H{r}<=6),\"Sim\",\"Não\")", "texto (fórmula)", "Sim quando atendida em até 6 h", "")])
T["fCustoDiaria"] = (fCustoDiaria, [
    ("id_unidade", None, None, "texto", "Unidade", REL_U),
    ("id_componente", None, None, "texto", "Componente da diária", "dComponenteCusto[id_componente]"),
    ("valor_diaria", R, None, "moeda", "Valor do componente na diária padrão", ""),
    ("fator_dia_de_espera", P, "=INDEX(dComponenteCusto!$C:$C,MATCH(B{r},dComponenteCusto!$A:$A,0))", "percentual (fórmula)", "Fator do componente que continua no dia de espera", ""),
    ("valor_dia_espera", R, "=C{r}*D{r}", "moeda (fórmula)", "Valor do componente por dia de espera", "")])
n_cd = len(fCensoDiario) + 1; n_cu = len(fCustoDiaria) + 1
T["fAtrasosDesocupacao"] = (fAtrasosDesocupacao, [
    ("id_atraso", None, None, "texto", "Código do episódio (chave)", ""),
    ("id_leito", None, None, "texto", "Leito ocupado durante a espera", "dLeito[id_leito]"),
    ("id_unidade", None, None, "texto", "Unidade", REL_U),
    ("prontuario", None, None, "texto", "Prontuário fictício", ""),
    ("id_motivo", None, None, "texto", "Motivo do atraso", "dMotivoDesocupacao[id_motivo]"),
    ("data_alta_medica", DTH, None, "data e hora", "Alta médica (início da espera)", ""),
    ("data_desocupacao", DTH, None, "data e hora", "Saída do leito (vazio se em aberto)", ""),
    ("situacao", None, None, "texto", "Resolvido ou Em aberto", ""),
    ("dias_atraso", F1, None, "decimal", "Dias entre a alta médica e a saída (em aberto: até o censo)", ""),
    ("custo_dia_espera", R, f"=ROUND(SUMIF(fCustoDiaria!$A$2:$A${n_cu},C{{r}},fCustoDiaria!$E$2:$E${n_cu}),2)", "moeda (fórmula)", "Custo de um dia de espera na unidade (soma de valor_dia_espera)", ""),
    ("custo_total", RS0, "=ROUND(I{r}*J{r},0)", "moeda (fórmula)", "Custo do episódio = dias de atraso × custo do dia de espera", ""),
    ("internacoes_perdidas", F2, f"=ROUND(I{{r}}/ROUND(SUMIF(fCensoDiario!$B$2:$B${n_cd},C{{r}},fCensoDiario!$E$2:$E${n_cd})/SUMIF(fCensoDiario!$B$2:$B${n_cd},C{{r}},fCensoDiario!$K$2:$K${n_cd}),1),2)",
     "decimal (fórmula)", "Leitos-dia perdidos ÷ permanência média da unidade nos 30 dias", ""),
    ("receita_nao_realizada", RS0, f"=ROUND(I{{r}}/ROUND(SUMIF(fCensoDiario!$B$2:$B${n_cd},C{{r}},fCensoDiario!$E$2:$E${n_cd})/SUMIF(fCensoDiario!$B$2:$B${n_cd},C{{r}},fCensoDiario!$K$2:$K${n_cd}),1)*INDEX(dUnidade!$H:$H,MATCH(C{{r}},dUnidade!$A:$A,0)),0)",
     "moeda (fórmula)", "Internações perdidas × receita média de AIH da unidade", "")])
T["dMetas"] = (dMetas, [
    ("indicador", None, None, "texto", "Indicador estratégico", ""),
    ("meta", "General", None, "número", "Meta ou limite de referência (percentuais em fração)", ""),
    ("unidade_medida", None, None, "texto", "Unidade de medida da meta", ""),
    ("sentido", None, None, "texto", "Menor é melhor, maior é melhor ou faixa", ""),
    ("regra", None, None, "texto", "Regra de situação usada no painel", "")])

ORDEM_ABAS = ["dUnidade", "dLeito", "dCalendario", "dProcedimento", "dMotivoDesocupacao", "dComponenteCusto",
              "fCensoDiario", "fLeitosCenso", "fSaidas", "fSolicitacoes", "fCustoDiaria", "fAtrasosDesocupacao", "dMetas"]
GRAN = {"dUnidade": "unidade de internação", "dLeito": "leito", "dCalendario": "dia (últimos 30 dias)", "dProcedimento": "procedimento × unidade",
        "dMotivoDesocupacao": "motivo de atraso na desocupação", "dComponenteCusto": "componente da diária", "fCensoDiario": "unidade × dia",
        "fLeitosCenso": "leito no censo de 19/10/2026 às 7h", "fSaidas": "saída de paciente", "fSolicitacoes": "pedido de leito ao NIR",
        "fCustoDiaria": "unidade × componente da diária", "fAtrasosDesocupacao": "episódio de atraso na desocupação", "dMetas": "indicador estratégico"}

# ---------------- planilha ----------------
ARIAL = Font(name="Arial", size=10); ARIAL_B = Font(name="Arial", size=10, bold=True)
TIT = Font(name="Arial", size=14, bold=True, color="1E3A5F"); AZUL = Font(name="Arial", size=10, color="0000FF")
CAB = PatternFill("solid", fgColor="DCE6F1"); AMARELO = PatternFill("solid", fgColor="FFFF00")
wb = Workbook(); leia = wb.active; leia.title = "LEIA-ME"
dicionario = []

def tabela(nome, linhas, cols):
    ws = wb.create_sheet(nome)
    for j, c in enumerate(cols, 1):
        cel = ws.cell(row=1, column=j, value=c[0]); cel.font = ARIAL_B
    for i, lin in enumerate(linhas, 2):
        for j, (col, fmt, formula, *_r) in enumerate(cols, 1):
            v = formula.replace("{r}", str(i)) if formula else lin.get(col)
            if v == "" and not formula: v = None
            cel = ws.cell(row=i, column=j, value=v); cel.font = ARIAL
            if fmt: cel.number_format = fmt
    ref = f"A1:{get_column_letter(len(cols))}{len(linhas) + 1}"
    t = Table(displayName=nome, ref=ref); t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = "A2"
    fixas = {"acao_recomendada": 80, "regra": 60, "procedimento": 44, "motivo": 40, "componente": 52, "pendencia": 40}
    for j, (col, *_r) in enumerate(cols, 1):
        if col in fixas: w = fixas[col]
        else:
            amostra = [len(str(l.get(col) or "")) for l in linhas[:200]]
            w = max(10, min(48, max([len(col)] + amostra) + 2))
            if cols[j - 1][1] in (DTH,): w = max(w, 17)
        ws.column_dimensions[get_column_letter(j)].width = w
    for col, fmt, formula, tipo, desc, rel in cols:
        dicionario.append({"tabela": nome, "coluna": col, "tipo": tipo, "descricao": desc, "relacionamento": rel})
    return ws

for nome in ORDEM_ABAS:
    tabela(nome, *T[nome])

# LEIA-ME
leia["A1"] = "Base de dados — Painel de Gestão de Leitos · Hospital Regional Horrara"; leia["A1"].font = TIT
leia["A2"] = "Dados fictícios para demonstração. Nenhum paciente, profissional ou valor é real."; leia["A2"].font = ARIAL
leia["A4"] = "Data de referência (censo)"; leia["A4"].font = ARIAL_B
leia["B4"] = HOJE; leia["B4"].number_format = DTH; leia["B4"].font = AZUL; leia["B4"].fill = AMARELO
leia["C4"] = "";
textos = [
    ("Período de análise", "20/09/2026 a 19/10/2026 (30 dias). O retrato dos leitos (fLeitosCenso) é o censo de 19/10/2026 às 7h."),
    ("Modelo", "Modelo estrela: tabelas com prefixo d são dimensões (cadastros) e com prefixo f são fatos (medições)."),
    ("Chave principal", "id_unidade (sigla da unidade) liga quase todas as tabelas; id_leito, id_procedimento, id_motivo e id_componente completam o modelo."),
    ("Percentuais", "Gravados em fração (0,85 = 85%), com formato de porcentagem."),
    ("Cores", "Azul: valor digitado que pode ser alterado. Fundo amarelo: data de referência usada nas fórmulas. Colunas com fórmula estão indicadas no Dicionario."),
    ("Recálculo", "As fórmulas são recalculadas ao abrir o arquivo. Altere B4 para simular outro momento do censo."),
    ("Conferência", "A aba Conferencia recalcula os principais indicadores a partir das tabelas e compara com os valores exibidos no painel."),
]
lin = 6
for k, v in textos:
    leia.cell(row=lin, column=1, value=k).font = ARIAL_B; c = leia.cell(row=lin, column=2, value=v); c.font = ARIAL; c.alignment = Alignment(wrap_text=True); lin += 1
lin += 1
for j, h in enumerate(["Aba", "Granularidade", "Linhas"], 1):
    c = leia.cell(row=lin, column=j, value=h); c.font = ARIAL_B; c.fill = CAB
for nome in ORDEM_ABAS:
    lin += 1
    leia.cell(row=lin, column=1, value=nome).font = ARIAL; leia.cell(row=lin, column=2, value=GRAN[nome]).font = ARIAL
    c = leia.cell(row=lin, column=3, value=len(T[nome][0])); c.font = ARIAL; c.number_format = I
leia.column_dimensions["A"].width = 24; leia.column_dimensions["B"].width = 90; leia.column_dimensions["C"].width = 10

# Conferencia
cf = wb.create_sheet("Conferencia")
kp = CALC["kpi_periodo"][-1]; ca = CALC["censo_atual"][-1]; dm = CALC["demanda"][-1]; du = CALC["desoc_unidade"][-1]
nl = len(fLeitosCenso) + 1; ns = len(fSolicitacoes) + 1; na = len(fAtrasosDesocupacao) + 1
CD = lambda c: f"fCensoDiario!${c}$2:${c}${n_cd}"
LC = lambda c: f"fLeitosCenso!${c}$2:${c}${nl}"
INDS = [  # indicador, fórmula na base, valor exibido no painel, formato, tolerância (precisão de exibição)
    ("Taxa de ocupação (30 dias)", f"=SUM({CD('E')})/SUM({CD('C')})", kp["taxa_ocupacao"] / 100, P, 0.0005),
    ("Ocupação instantânea (censo)", f"=COUNTIF({LC('C')},\"Ocupado\")/(COUNTA({LC('A')})-COUNTIF({LC('C')},\"Bloqueado\"))", ca["taxa_ocupacao"] / 100, P, 0.0005),
    ("Média de permanência (dias)", f"=SUM({CD('E')})/SUM({CD('K')})", kp["media_permanencia"], F1, 0.05),
    ("Índice de giro (saídas por leito)", f"=SUM({CD('K')})/(SUM({CD('C')})/30)", kp["giro"], F1, 0.05),
    ("Mortalidade institucional", f"=(SUM({CD('H')})-SUM({CD('I')}))/SUM({CD('K')})", kp["mortalidade_institucional"] / 100, P, 0.0005),
    ("Leitos vagos (censo)", f"=COUNTIF({LC('C')},\"Vago\")", ca["vagos"], I, 0),
    ("Pacientes em vermelho (Kanban)", f"=COUNTIF({LC('N')},\"Vermelho\")", ca["vermelhos"], I, 0),
    ("Pacientes de alta aguardando desocupação", f"=COUNTIF({LC('D')},\"Aguardando desocupação\")", ca["aguardando_desocupacao"], I, 0),
    ("Solicitações pendentes no NIR", f"=COUNTIF(fSolicitacoes!$G$2:$G${ns},\"Pendente\")", dm["pendentes"], I, 0),
    ("Leitos-dia perdidos com desocupação", f"=SUM(fAtrasosDesocupacao!$I$2:$I${na})", du["dias_perdidos"], F1, 0.05),
    ("Custo da espera nos 30 dias (R$)", f"=SUM(fAtrasosDesocupacao!$K$2:$K${na})", du["custo_total"], RS0, 0.5),
]
cf["A1"] = "Conferência dos indicadores: base × painel"; cf["A1"].font = TIT
for j, h in enumerate(["Indicador", "Como é calculado na base", "Calculado na base", "Valor no painel", "Diferença", "Tolerância"], 1):
    c = cf.cell(row=3, column=j, value=h); c.font = ARIAL_B; c.fill = CAB
COMO = ["Σ pacientes-dia ÷ Σ leitos operacionais", "Ocupados ÷ (leitos − bloqueados)", "Σ pacientes-dia ÷ Σ saídas", "Σ saídas ÷ leitos operacionais médios",
        "(óbitos − óbitos até 24 h) ÷ saídas", "Contagem de leitos vagos", "Contagem de classificação Vermelho", "Contagem da etapa Aguardando desocupação",
        "Contagem de pedidos pendentes", "Σ dias de atraso", "Σ custo dos episódios"]
for i, (nm, f, v, fmt, tol) in enumerate(INDS, 4):
    cf.cell(row=i, column=1, value=nm).font = ARIAL; cf.cell(row=i, column=2, value=COMO[i - 4]).font = ARIAL
    c = cf.cell(row=i, column=3, value=f); c.font = ARIAL; c.number_format = fmt
    c = cf.cell(row=i, column=4, value=v); c.font = AZUL; c.number_format = fmt
    c = cf.cell(row=i, column=5, value=f"=C{i}-D{i}"); c.font = ARIAL; c.number_format = "0.0000"
    c = cf.cell(row=i, column=6, value=tol); c.font = AZUL; c.number_format = "0.0000"
fim = 3 + len(INDS)
cf.cell(row=fim + 2, column=1, value="Resultado").font = ARIAL_B
cf.cell(row=fim + 2, column=2, value=f"=IF(SUMPRODUCT((ABS(E4:E{fim})>F4:F{fim})*1)=0,\"Base confere com o painel\",\"Revisar: há diferença acima da tolerância\")").font = ARIAL_B
cf.cell(row=fim + 3, column=1, value="Nota").font = ARIAL
cf.cell(row=fim + 3, column=2, value="O painel exibe valores arredondados; a tolerância de cada linha corresponde à precisão exibida.").font = ARIAL
for col, w in zip("ABCDEF", [40, 46, 20, 18, 14, 12]): cf.column_dimensions[col].width = w

# Dicionario
wsd = wb.create_sheet("Dicionario")
for j, h in enumerate(["tabela", "coluna", "tipo", "descricao", "relacionamento"], 1): wsd.cell(row=1, column=j, value=h).font = ARIAL_B
for i, r in enumerate(dicionario, 2):
    for j, k in enumerate(["tabela", "coluna", "tipo", "descricao", "relacionamento"], 1): wsd.cell(row=i, column=j, value=r[k]).font = ARIAL
t = Table(displayName="Dicionario", ref=f"A1:E{len(dicionario) + 1}"); t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
wsd.add_table(t); wsd.freeze_panes = "A2"
for col, w in zip("ABCDE", [22, 30, 20, 80, 30]): wsd.column_dimensions[col].width = w

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(f"{SAIDA}/base-dados-leitos.xlsx")

# ---------------- CSV (um por tabela, valores calculados) e JSON ----------------
def fmt_csv(v):
    if v is None: return ""
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date): return v.isoformat()
    if isinstance(v, float): return f"{v:.6f}".rstrip("0").rstrip(".")
    return v
for nome in ORDEM_ABAS:
    linhas, cols = T[nome]
    with open(f"{SAIDA}/csv/{nome}.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f); w.writerow([c[0] for c in cols])
        for l in linhas: w.writerow([fmt_csv(l.get(c[0])) for c in cols])
with open(f"{SAIDA}/csv/Dicionario.csv", "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f); w.writerow(["tabela", "coluna", "tipo", "descricao", "relacionamento"])
    for r in dicionario: w.writerow([r[k] for k in ["tabela", "coluna", "tipo", "descricao", "relacionamento"]])
json.dump({"hoje": HOJE.strftime("%Y-%m-%dT%H:%M"), "hospital": "Hospital Regional Horrara", "observacao": "Dados fictícios",
           "tabelas": {n: [{k: fmt_csv(v) for k, v in l.items()} for l in T[n][0]] for n in ORDEM_ABAS}},
          open(f"{SAIDA}/dados.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print({n: len(T[n][0]) for n in ORDEM_ABAS}, "dicionário:", len(dicionario))
